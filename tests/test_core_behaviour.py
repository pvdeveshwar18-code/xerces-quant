import importlib
import io
import os
import sys
import tempfile
import types
import unittest

import pandas as pd
from openpyxl import load_workbook


def install_runtime_stubs():
    streamlit = types.ModuleType("streamlit")
    streamlit.secrets = {}
    streamlit.session_state = {}

    def cache_data(*args, **kwargs):
        def decorator(func):
            return func
        return decorator

    streamlit.cache_data = cache_data
    sys.modules["streamlit"] = streamlit

    yfinance = types.ModuleType("yfinance")
    yfinance.download = lambda *args, **kwargs: pd.DataFrame()
    sys.modules["yfinance"] = yfinance

    plotly = types.ModuleType("plotly")
    graph_objects = types.ModuleType("plotly.graph_objects")

    class Trace:
        def __init__(self, *args, **kwargs):
            self.args = args
            self.kwargs = kwargs

    class Figure:
        def __init__(self, *args, **kwargs):
            self.args = args
            self.kwargs = kwargs
            self.traces = []

        def add_trace(self, trace):
            self.traces.append(trace)

        def update_layout(self, *args, **kwargs):
            self.layout_args = args
            self.layout_kwargs = kwargs

    graph_objects.Figure = Figure
    graph_objects.Treemap = Trace
    graph_objects.Scatter = Trace
    graph_objects.Heatmap = Trace
    sys.modules["plotly"] = plotly
    sys.modules["plotly.graph_objects"] = graph_objects


class XercesPlusCoreTests(unittest.TestCase):
    def setUp(self):
        install_runtime_stubs()
        self.tempdir = tempfile.TemporaryDirectory()
        os.environ["XERCES_DATA_DIR"] = self.tempdir.name
        os.environ.pop("EMERGENT_LLM_KEY", None)
        sys.modules.pop("xerces_plus", None)
        self.xp = importlib.import_module("xerces_plus")

    def tearDown(self):
        self.tempdir.cleanup()
        os.environ.pop("XERCES_DATA_DIR", None)
        os.environ.pop("EMERGENT_LLM_KEY", None)

    def test_llm_key_has_no_hardcoded_fallback(self):
        self.assertEqual(self.xp._get_llm_key(), "")
        os.environ["EMERGENT_LLM_KEY"] = "test-key"
        self.assertEqual(self.xp._get_llm_key(), "test-key")

    def test_watchlist_persists_and_deduplicates(self):
        self.assertTrue(self.xp.add_to_watchlist("RELIANCE.NS", "Reliance"))
        self.assertFalse(self.xp.add_to_watchlist("RELIANCE.NS", "Reliance"))
        self.assertEqual(self.xp.load_watchlist(), [{
            "ticker": "RELIANCE.NS",
            "name": "Reliance",
            "added": self.xp.load_watchlist()[0]["added"],
        }])

    def test_excel_export_writes_fallback_sheet(self):
        workbook_bytes = self.xp.build_excel_bytes({})
        workbook = load_workbook(io.BytesIO(workbook_bytes))
        self.assertIn("Summary", workbook.sheetnames)

    def test_journal_analysis_frame_renames_numeric_columns(self):
        journal = pd.DataFrame([{
            "Date": "2026-08-08",
            "Ticker": "TCS.NS",
            "Side": "Long",
            "Qty": "2",
            "Entry": "100",
            "Exit": "110",
            self.xp.JOURNAL_COLS[6]: "20",
            "P&L %": "10",
            "Strategy": "Breakout",
            "Notes": "",
        }])
        analysis = self.xp._journal_analysis_frame(journal)
        self.assertIn("P&L", analysis.columns)
        self.assertEqual(float(analysis.loc[0, "P&L"]), 20.0)
        self.assertEqual(float(analysis.loc[0, "Entry_Price"]), 100.0)


if __name__ == "__main__":
    unittest.main()
